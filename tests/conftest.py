import json

import pymupdf as fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MAIC_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("MAIC_DATABASE_URL", "sqlite:///" + str(tmp_path / "app.db"))
    monkeypatch.setenv("MAIC_CHECKPOINT_URL", "sqlite:///" + str(tmp_path / "checkpoint.db"))
    monkeypatch.setenv("MAIC_MODE", "demo")
    monkeypatch.setenv("MAIC_STORAGE", "local")
    monkeypatch.setenv("MAIC_MAX_AUTO_ROUNDS", "0")
    monkeypatch.setenv("MAIC_MAX_MODEL_CALLS", "12")
    monkeypatch.setenv("MAIC_CONTEXT_TOKENS", "4000")
    monkeypatch.setenv("MAIC_SCRIPT_BATCH_SIZE", "8")
    from app.config import settings
    from app.db import engine
    settings.cache_clear()
    engine.cache_clear()
    from app.main import app
    with TestClient(app) as test_client:
        yield test_client
    engine().dispose()
    engine.cache_clear()
    settings.cache_clear()


@pytest.fixture
def pdf_bytes():
    doc = fitz.open()
    for heading, content in [
        ("Evaporation", "The Sun provides energy. Liquid water becomes water vapour when it evaporates."),
        ("Condensation", "Water vapour cools and forms droplets. Water droplets form clouds."),
        ("Precipitation", "Water returns to the ground as rain or snow. The water cycle continues.")]:
        page = doc.new_page()
        page.insert_text((50, 60), heading, fontsize=22)
        page.insert_textbox(fitz.Rect(50, 90, 540, 400), content, fontsize=14)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def fake_render(monkeypatch):
    """Workflow tests isolate LibreOffice; test the real renderer separately."""
    from app.schemas import digest
    from app.storage import Storage
    monkeypatch.setattr("app.workflow.renderer_backend", lambda: "test")
    def render(job_id, draft, source, mode, media=None):
        storage = Storage()
        prefix = f"jobs/{job_id}/test-render/{digest(draft)}"
        out = storage.path(prefix)
        out.mkdir(parents=True, exist_ok=True)
        prs = Presentation()
        pdf = fitz.open()
        files = {}
        for i, slide in enumerate(draft.slide_content.slides):
            prs.slides.add_slide(prs.slide_layouts[6])
            pdf.new_page().insert_text((30, 40), slide.id)
            path = out / "render" / f"slide-{i+1}.png"
            path.parent.mkdir(exist_ok=True)
            Image.new("RGB", (960, 540), "white").save(path)
            files[f"render/slide-{i+1}.png"] = f"{prefix}/render/slide-{i+1}.png"
        prs.save(out / "slides.pptx")
        pdf.save(out / "slides.pdf")
        pdf.close()
        for name, data in [("geometry.json", {"geometry": [], "hard_count": 0}), ("lint.json", {"findings": []})]:
            storage.put_json(f"{prefix}/{name}", data)
        files.update({name: f"{prefix}/{name}" for name in ["slides.pptx", "slides.pdf", "geometry.json", "lint.json"]})
        return {"prefix": prefix, "files": files, "slide_count": len(prs.slides)}
    def convert(path):
        pdf = fitz.open()
        pdf.new_page().insert_text((30, 50), "Test lesson plan")
        pdf.save(path.with_suffix(".pdf"))
        pdf.close()
        return path.with_suffix(".pdf")
    monkeypatch.setattr("app.workflow.render_draft", render)
    monkeypatch.setattr("app.workflow.convert_docx_pdf", convert)


def upload(client, pdf_bytes, subject="Science"):
    result = client.post("/api/jobs", files={"file": ("lesson.pdf", pdf_bytes, "application/pdf")},
                         data={"request": json.dumps({"subject": subject, "audience": "Grade 6", "duration_minutes": 30,
                                                      "slide_count": 3})})
    assert result.status_code == 201, result.text
    return result.json()["id"]
