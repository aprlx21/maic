import copy
import io

from pptx import Presentation

from app.worker import run_once
from tests.conftest import upload


def current(client, job_id):
    detail = client.get(f"/api/jobs/{job_id}").json()
    assert detail["status"] != "failed", detail["error"]
    return detail


def decision(client, job_id, revision, action="approve", **extra):
    return client.post(f"/api/jobs/{job_id}/decisions", json={
        "action": action, "revision_id": revision["id"], "approval_hash": revision["approval_hash"],
        "teacher_name": "Teacher test", **extra})


def test_review_edit_approve_and_immutable_versions(client, pdf_bytes, fake_render):
    job_id = upload(client, pdf_bytes)
    assert run_once()
    first = current(client, job_id)
    assert first["status"] == "awaiting_teacher"
    old = first["revision"]
    assert {r["peer"] for r in old["reviews"]} == {
        "pedagogy", "content", "assessment", "slide_critic", "validation", "render"}
    draft = copy.deepcopy(old["draft"])
    draft["slide_content"]["slides"][0]["teaching_script"]["narration"] = "Teacher edited narration."
    assert decision(client, job_id, old, "edit", draft=draft).status_code == 202
    assert run_once()
    edited = current(client, job_id)["revision"]
    assert edited["id"] != old["id"]
    assert edited["approval_hash"] != old["approval_hash"]
    assert decision(client, job_id, old).status_code == 409
    assert decision(client, job_id, edited).status_code == 202
    assert run_once()
    published = current(client, job_id)
    assert published["status"] == "published"
    publication = published["publications"][0]
    assert "speaker-notes.pdf" in publication["artifacts"]
    data = client.get(f"/api/jobs/{job_id}/asset", params={"key": publication["artifacts"]["slides.pptx"]}).content
    pptx = Presentation(io.BytesIO(data))
    assert "Teacher edited narration." in pptx.slides[0].notes_slide.notes_text_frame.text
    package_v1 = client.get(f"/api/lessons/{job_id}/versions/1").json()
    assert package_v1["approval_hash"] == edited["approval_hash"]
    assert {a["type"] for a in package_v1["teaching_actions"]} >= {"ShowFile", "ReadScript"}
    draft["lesson"]["title"] = "New title for version two"
    assert decision(client, job_id, edited, "edit", draft=draft).status_code == 202
    assert run_once()
    latest = current(client, job_id)["revision"]
    assert decision(client, job_id, latest).status_code == 202
    assert run_once()
    assert len(current(client, job_id)["publications"]) == 2
    assert client.get(f"/api/lessons/{job_id}/versions/1").json() == package_v1


def test_job_keeps_source_snapshot_after_ocr_config_change(client, pdf_bytes, fake_render, tmp_path):
    from app.config import settings
    job_id = upload(client, pdf_bytes)
    run_once()
    before = current(client, job_id)["source"]
    settings().ocr_languages = "eng"
    settings().tessdata_prefix = str(tmp_path / "different-ocr")
    after = current(client, job_id)["source"]
    assert before == after
    key = before["pages"][0]["preview_key"]
    assert client.get(f"/api/jobs/{job_id}/asset", params={"key": key}).status_code == 200


def test_invalid_evidence_blocks_and_limits_auto_repair(client, pdf_bytes, fake_render):
    from app.config import settings
    settings().max_auto_rounds = 2
    job_id = upload(client, pdf_bytes)
    run_once()
    initial = current(client, job_id)["revision"]
    draft = copy.deepcopy(initial["draft"])
    draft["slide_content"]["slides"][0]["evidence_ids"] = ["invented-source"]
    assert decision(client, job_id, initial, "edit", draft=draft).status_code == 202
    run_once()
    detail = current(client, job_id)
    assert detail["status"] == "awaiting_teacher"
    assert len(detail["revisions"]) == 4  # Initial + edited + two automatic rounds.
    assert any(f["severity"] == "blocked" for r in detail["revision"]["reviews"] for f in r["findings"])
    assert decision(client, job_id, detail["revision"]).status_code == 409
    assert detail["publications"] == []


def test_revise_findings_require_recorded_reason(client, pdf_bytes, fake_render):
    job_id = upload(client, pdf_bytes)
    run_once()
    initial = current(client, job_id)["revision"]
    draft = copy.deepcopy(initial["draft"])
    draft["lesson"]["activities"][0]["minutes"] += 2
    assert decision(client, job_id, initial, "edit", draft=draft).status_code == 202
    run_once()
    revision = current(client, job_id)["revision"]
    assert decision(client, job_id, revision).status_code == 409
    reasons = {f["id"]: "Teacher will adjust pacing in class."
               for r in revision["reviews"] for f in r["findings"] if f["severity"] == "revise"}
    assert reasons
    assert decision(client, job_id, revision, accepted_findings=reasons).status_code == 202
    run_once()
    assert current(client, job_id)["status"] == "published"


def test_upload_and_asset_guards(client, pdf_bytes):
    result = client.post("/api/jobs", files={"file": ("fake.pdf", b"not a PDF", "application/pdf")},
                         data={"request": '{"subject":"Science","audience":"Grade 6","duration_minutes":30}'})
    assert result.status_code == 422
    job_id = upload(client, pdf_bytes)
    assert client.get(f"/api/jobs/{job_id}/asset", params={"key": "../../.env"}).status_code == 404
    assert client.get(f"/api/lessons/{job_id}/versions/1").status_code == 404


def test_changed_artifact_cannot_be_published(client, pdf_bytes, fake_render):
    from app.storage import Storage
    job_id = upload(client, pdf_bytes)
    run_once()
    revision = current(client, job_id)["revision"]
    Storage().put(revision["artifacts"]["files"]["slides.pdf"], b"changed after review")
    assert decision(client, job_id, revision).status_code == 202
    run_once()
    detail = client.get(f"/api/jobs/{job_id}").json()
    assert detail["status"] == "failed"
    assert detail["publications"] == []


def test_retry_resumes_checkpoint_without_rereading(client, pdf_bytes, fake_render, monkeypatch):
    from app import workflow
    job_id = upload(client, pdf_bytes)
    original = workflow.create_draft
    def fail(*args, **kwargs):
        raise RuntimeError("Temporary generation failure")
    monkeypatch.setattr(workflow, "create_draft", fail)
    run_once()
    assert client.get(f"/api/jobs/{job_id}").json()["status"] == "failed"
    monkeypatch.setattr(workflow, "create_draft", original)
    assert client.post(f"/api/jobs/{job_id}/retry").status_code == 202
    run_once()
    detail = current(client, job_id)
    assert detail["status"] == "awaiting_teacher"
    assert len([e for e in detail["events"] if e["details"].get("pages")]) == 1
