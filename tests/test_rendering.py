import shutil

import pymupdf as fitz
import pytest
from pptx import Presentation

from app.rendering import build_deck, convert_docx_pdf, renderer_backend, upstream_modules
from app.schemas import Draft, SourceDocument
from app.worker import run_once
from tests.conftest import upload


def draft_from_job(client, pdf_bytes):
    job_id = upload(client, pdf_bytes)
    run_once()
    detail = client.get(f"/api/jobs/{job_id}").json()
    assert detail["status"] == "awaiting_teacher", detail["error"]
    return Draft.model_validate(detail["revision"]["draft"]), SourceDocument.model_validate(detail["source"])


@pytest.mark.parametrize("layout", ["points", "statement", "example", "two_columns"])
def test_real_deckkit_build_and_static_lint(client, pdf_bytes, fake_render, tmp_path, monkeypatch, layout):
    """Exercise upstream DeckKit and lint; isolate only the unavailable PDF conversion process."""
    draft, source = draft_from_job(client, pdf_bytes)
    for design in draft.design.designs:
        design.layout = layout
    if layout == "example":
        for slide in draft.slide_content.slides:
            slide.practical_example = "Tình huống giả định: Quan sát cốc nước lạnh và giải thích các giọt nước bên ngoài."
    _, lint, render = upstream_modules()
    def converter(soffice, src, dest):
        doc = fitz.open()
        for _slide in Presentation(src).slides:
            page = doc.new_page(width=960, height=540)
            page.insert_text((40, 50), "Layout test (PDF converter isolated)")
        doc.save(dest)
        doc.close()
        return dest
    monkeypatch.setattr(render, "render_accessible_pdf", converter)
    monkeypatch.setattr("app.rendering.soffice_path", lambda: "test-converter")
    # Arial is available on Windows; production Docker supplies Noto Sans.
    from app.config import settings
    settings().slide_font = "Arial"
    result = build_deck(draft, source, tmp_path / "deckkit", "demo")
    assert not [f for f in result["geometry"] if f[1] == "CRITICAL"], result
    assert result["hard_count"] > 0  # Real pixel lint correctly rejects the intentionally mismatched previews.
    (tmp_path / "static-only").mkdir()
    shutil.copyfile(tmp_path / "deckkit/slides.pptx", tmp_path / "static-only/slides.pptx")
    assert lint.lint(str(tmp_path / "static-only/slides.pptx"), static_ok=True) == 0
    pptx = Presentation(tmp_path / "deckkit/slides.pptx")
    assert len(pptx.slides) == 3
    assert draft.slide_content.slides[0].teaching_script.narration in pptx.slides[0].notes_slide.notes_text_frame.text


def test_real_pdf_render(client, pdf_bytes, fake_render, tmp_path):
    try:
        renderer_backend()
    except RuntimeError:
        pytest.skip("Requires LibreOffice or Microsoft Office")
    from app.config import settings
    if renderer_backend() == "microsoft-office":
        settings().slide_font = "Arial"
    draft, source = draft_from_job(client, pdf_bytes)
    result = build_deck(draft, source, tmp_path / "real", "demo")
    assert result["hard_count"] == 0
    with fitz.open(tmp_path / "real/slides.pdf") as doc:
        assert len(doc) == len(draft.slide_content.slides)
        assert all(page.get_text().strip() for page in doc)
    from docx import Document
    lesson = Document()
    lesson.add_paragraph("MAIC lesson export")
    lesson.save(tmp_path / "lesson.docx")
    with fitz.open(convert_docx_pdf(tmp_path / "lesson.docx")) as doc:
        assert "MAIC lesson export" in doc[0].get_text()
