"""Controlled slides_maker integration: JSON in, native editable PPTX + rendered pages out."""
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from functools import lru_cache
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.shared import Inches as DocInches
from docx.shared import Pt

from app.budget import remaining_seconds
from app.config import settings
from app.schemas import Draft, SourceDocument, digest
from app.storage import Storage


def soffice_path():
    cfg = settings()
    choices = [cfg.soffice, shutil.which("soffice"),
               r"C:\Program Files\LibreOffice\program\soffice.exe"]
    for choice in choices:
        if choice and Path(choice).exists():
            return str(choice)
    raise RuntimeError("Cần LibreOffice để render và xuất PDF. Cài LibreOffice hoặc chạy bằng Docker Compose.")


def renderer_backend():
    try:
        soffice_path()
        return "libreoffice"
    except RuntimeError:
        if sys.platform == "win32":
            import winreg
            try:
                for name in ("PowerPoint.Application", "Word.Application"):
                    with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, name + r"\CLSID"):
                        pass
                return "microsoft-office"
            except OSError:
                pass
        raise


def office_pdf(source: Path, destination: Path):
    """Local Windows demo fallback; production workers use LibreOffice."""
    from filelock import FileLock
    env = os.environ.copy()
    env.pop("OPENAI_API_KEY", None)
    script = Path(__file__).resolve().parents[1] / "scripts/office_pdf.ps1"
    with FileLock(str(settings().data_dir / "office-render.lock"), timeout=180):
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-NonInteractive", "-File", str(script),
             "-Kind", source.suffix[1:], "-Source", str(source.resolve()),
             "-Destination", str(destination.resolve())],
            capture_output=True, timeout=180, env=env, check=False)
    if result.returncode or not destination.exists():
        raise RuntimeError("Microsoft Office không xuất được PDF; kiểm tra Office đã được kích hoạt.")
    return destination


@lru_cache
def upstream_modules():
    os.environ["SLIDE_MAKER_NO_ENV_CHECK"] = "1"
    os.environ["SLIDE_MAKER_NO_VERSION_CHECK"] = "1"
    path = settings().vendor / "scripts"
    if not (path / "deckkit.py").exists():
        raise RuntimeError("Chưa có slides_maker. Chạy python scripts/bootstrap_vendor.py")
    repo = settings().vendor.parents[1]
    expected = json.loads((repo.parent / "slides_maker.lock.json").read_text())["commit"]
    actual = subprocess.run(["git", "-C", str(repo), "rev-parse", "HEAD"], capture_output=True,
                            text=True, check=True).stdout.strip()
    if actual != expected:
        raise RuntimeError("Commit slides_maker không khớp lock; chạy scripts/bootstrap_vendor.py trước khi khởi động worker.")
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))
    import deckkit
    import lint_deck
    import render_deck
    return deckkit, lint_deck, render_deck


def script_notes(slide):
    script = slide.teaching_script
    if not script:
        return ""
    parts = [script.introduction, script.narration, script.transition]
    for question in script.questions:
        parts += [f"HỎI ({question.wait_seconds}s): {question.prompt}",
                  f"ĐÁP ÁN DÀNH CHO GIÁO VIÊN: {question.expected_answer}"]
    parts += [f"Thời lượng giảng dự kiến: {script.estimated_seconds}s", "Nguồn: " + ", ".join(script.evidence_ids)]
    return "\n\n".join(p for p in parts if p)


def image_credit(media):
    if not media.get("asset_key"):
        return ""
    if media["origin"] == "ai":
        return "Minh họa AI · không phải ảnh sự kiện thật"
    return f"{media['title']} — {media['author']} · {media['license']}\n{media.get('license_url', '')}\n{media.get('changes', '')}"


def build_deck(draft: Draft, source: SourceDocument, out: Path, mode: str, media=None):
    dk, lint, render = upstream_modules()
    from pptx.dml.color import RGBColor
    cfg = settings()
    storage = Storage()
    out.mkdir(parents=True, exist_ok=True)
    dk.FONT = cfg.slide_font
    dk.EAFONT = cfg.slide_font
    dk.DISPLAY = cfg.slide_font
    palette = {"ocean": ("153B50", "007C91"), "forest": ("183C30", "32806A"), "ink": ("252B3B", "7257A5")}
    dark, accent = [RGBColor.from_string(c) for c in palette[draft.design.theme]]
    text_color = RGBColor.from_string("253445")
    prs = dk.blank_deck(13.333, 7.5)
    designs = {d.slide_id: d for d in draft.design.designs}
    units = {u.id: u for u in source.units}
    media = media or {}
    (out / "image-credits.json").write_text(json.dumps(media, ensure_ascii=False, indent=2), encoding="utf-8")
    for index, spec in enumerate(draft.slide_content.slides):
        slide = dk.add_slide(prs)
        dk.slide_background(slide, RGBColor.from_string("FAFBFC"))
        dk.box(slide, .0, .0, 13.333, .12, fill=accent)
        def write(value, x, y, w, h, size=24, bold=False, color=text_color, slide=slide):
            fitted = dk.fit_text_size([(value, bold)], w, h, size, font=cfg.slide_font, min_size=18)
            return dk.text(slide, x, y, w, h, [[(value, fitted, color, bold, False, cfg.slide_font)]], space_after=0)
        write(spec.title, .6, .5, 12.1, 1.05, 31, True, dark)
        layout = designs[spec.id].layout if spec.id in designs else "points"
        if layout != "statement":
            write(spec.takeaway, .65, 1.75, 11.95, .7, 23, True, accent)
        bullets = spec.bullets
        visual = media.get(spec.id, {})
        image_path = None
        if visual.get("asset_key"):
            import hashlib
            raw = storage.get(visual["asset_key"])
            if hashlib.sha256(raw).hexdigest() != visual["sha256"]:
                raise ValueError("Ảnh minh họa đã thay đổi sau khi chuẩn bị")
            image_path = out / "images" / f"image-{index+1}.jpg"
            image_path.parent.mkdir(exist_ok=True)
            image_path.write_bytes(raw)
        example_focus = layout == "example" and bool(spec.practical_example)
        bottom = 5.6 if spec.practical_example and not example_focus else 6.3
        if layout == "statement":
            dk.box(slide, .65, 2.0, 5.6, bottom - 2.0, fill=dark)
            write(spec.takeaway, .95, 2.3, 5.0, bottom - 2.6, 32, True, RGBColor.from_string("FFFFFF"))
            height = (bottom - 2.0) / len(bullets)
            for j, value in enumerate(bullets):
                write(value, 6.75, 2.0 + j * height, 5.85, height - .12, 24)
        elif example_focus:
            dk.box(slide, .65, 2.7, 5.8, 3.6, fill=RGBColor.from_string("EAF2F5"))
            write("TÌNH HUỐNG VẬN DỤNG", .95, 2.9, 5.2, .5, 20, True, accent)
            write(spec.practical_example, .95, 3.6, 5.2, 2.3, 26, False, dark)
            height = 3.6 / len(bullets)
            for j, value in enumerate(bullets):
                write(value, 6.95, 2.7 + j * height, 5.65, height - .12, 23)
        elif layout == "figure" and spec.figure_id in units and units[spec.figure_id].asset_key:
            unit = units[spec.figure_id]
            storage.get(unit.asset_key)
            dk.picture(slide, str(storage.path(unit.asset_key)), 6.8, 2.65, 5.85, bottom - 2.65,
                       fit="contain", alt=unit.content)
            for j, value in enumerate(bullets):
                height = (bottom - 2.75) / len(bullets)
                write(value, .7, 2.65 + j * height, 5.55, height - .06, 22)
        elif layout == "photo" and image_path:
            # Whole image remains visible: science details and license attribution are not cropped away.
            dk.picture(slide, str(image_path), 5.7, 2.6, 6.95, bottom - 2.95,
                       fit="contain", alt=visual.get("caption") or spec.title)
            height = (bottom - 2.75) / len(bullets)
            for j, value in enumerate(bullets):
                write(value, .7, 2.65 + j * height, 4.55, height - .08, 22)
            caption = visual.get("caption") or "Quan sát hình minh họa"
            write(caption, 5.7, bottom - .3, 6.95, .42, 18, False, accent)
        elif layout == "process" and len(bullets) >= 2:
            count = len(bullets)
            gap = .38
            width = (11.95 - (count-1) * gap) / count
            for j, value in enumerate(bullets):
                x = .7 + j * (width + gap)
                dk.box(slide, x, 2.85, width, bottom - 3.0, fill=RGBColor.from_string("EAF2F5"))
                write(f"{j+1:02}", x + .18, 3.0, width - .36, .5, 25, True, accent)
                write(value, x + .18, 3.65, width - .36, bottom - 3.9, 22)
                if j < count-1:
                    dk.arrow(slide, x + width + .05, 3.5, .27, .18, color=accent)
        elif layout == "two_columns" and len(bullets) >= 2:
            mid = (len(bullets) + 1) // 2
            for col, values in enumerate((bullets[:mid], bullets[mid:])):
                dk.box(slide, .65 + col * 6.2, 2.65, 5.8, bottom - 2.65,
                       fill=RGBColor.from_string("EAF2F5" if col == 0 else "F0F2F5"))
                for j, value in enumerate(values):
                    height = (bottom - 2.8) / len(values)
                    write(value, .9 + col * 6.2, 2.8 + j * height, 5.3, height - .12, 24)
        else:
            for j, value in enumerate(bullets):
                height = (bottom - 2.8) / len(bullets)
                write(value, .75, 2.8 + j * height, 11.8, height - .1,
                      27 if layout == "statement" else 24)
        if spec.practical_example and not example_focus:
            dk.box(slide, .65, 5.85, 12.0, .85, fill=RGBColor.from_string("EAF2F5"))
            write(spec.practical_example, .85, 5.97, 11.6, .62, 19, False, dark)
        source_pages = sorted({units[e].page for e in [*spec.evidence_ids, spec.figure_id] if e in units})
        tag = ("DEMO · " if mode == "demo" else "") + "Nguồn PDF: trang " + ", ".join(map(str, source_pages))
        section = next((k for k in draft.slide_content.sections if spec.id in k.slide_ids), None)
        if section:
            tag = section.id.upper() + " · " + tag
        if visual.get("asset_key"):
            tag += " | " + ("Minh họa AI" if visual["origin"] == "ai" else "Ảnh Internet · " + visual["license"])
        dk.footer(slide, tag=tag, page=index + 1)
        slide.notes_slide.notes_text_frame.text = script_notes(spec) + "\n\n" + image_credit(visual) + "\n" + visual.get("source_url", "")
    pptx = out / "slides.pptx"
    prs.save(pptx)
    geometry = dk.lint_layout(prs, verbose=False, strict=False)
    # Keep factual intermediate artifacts, not synthetic claims of upstream human approval.
    (out / "content-plan.json").write_text(draft.slide_content.model_dump_json(indent=2), encoding="utf-8")
    (out / "design-plan.json").write_text(draft.design.model_dump_json(indent=2), encoding="utf-8")
    pdf = out / "slides.pdf"
    if renderer_backend() == "microsoft-office":
        made = office_pdf(pptx, pdf)
    else:
        with contextlib.redirect_stdout(io.StringIO()):
            made = render.render_accessible_pdf(soffice_path(), str(pptx), str(pdf))
    if not made:
        raise RuntimeError("LibreOffice không chuyển được PPTX thành PDF")
    preview_dir = out / "render"
    preview_dir.mkdir(exist_ok=True)
    with fitz.open(pdf) as rendered:
        if len(rendered) != len(draft.slide_content.slides):
            raise RuntimeError("Số trang PDF render không khớp số slide")
        for i, page in enumerate(rendered):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(preview_dir / f"slide-{i+1}.png")
    log = io.StringIO()
    with contextlib.redirect_stdout(log):
        hard_count = lint.lint(str(pptx), json_out=str(out / "lint.json"), renders_dir=str(preview_dir),
                               static_ok=True, notes_ok=not all(s.teaching_script for s in draft.slide_content.slides))
    (out / "lint.log").write_text(log.getvalue(), encoding="utf-8")
    return {"geometry": geometry, "hard_count": hard_count}


def render_draft(job_id, draft, source, mode, media=None):
    storage = Storage()
    visual = draft.model_dump()
    for slide in visual["slide_content"]["slides"]:
        slide["teaching_script"] = None
    backend = renderer_backend()
    media = media or {}
    render_hash = digest({"draft": visual, "font": settings().slide_font, "adapter": "v4-knowledge-layouts", "media": media,
                          "renderer": backend, "mode": mode})
    prefix = f"jobs/{job_id}/renders/{render_hash}"
    manifest_key = prefix + "/manifest.json"
    if storage.exists(manifest_key):
        return storage.get_json(manifest_key)
    out = settings().data_dir / "work" / job_id / render_hash
    out.mkdir(parents=True, exist_ok=True)
    spec = out / "input.json"
    spec.write_text(json.dumps({"draft": visual, "source": source.model_dump(), "mode": mode, "media": media}, ensure_ascii=False), encoding="utf-8")
    temp = out / "tmp"
    temp.mkdir(exist_ok=True)
    env = {**os.environ, "TMP": str(temp), "TEMP": str(temp), "TMPDIR": str(temp),
           "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1"}
    env.pop("OPENAI_API_KEY", None)
    result = subprocess.run([sys.executable, "-m", "app.rendering", str(spec)], capture_output=True,
                            text=True, encoding="utf-8", errors="replace", timeout=remaining_seconds(job_id),
                            check=False, env=env)
    if result.returncode:
        raise RuntimeError("Dựng slide thất bại: " + (result.stderr or result.stdout)[-2500:])
    artifacts = {}
    for path in out.rglob("*"):
        if path.is_file() and path.name != "input.json" and not path.is_relative_to(temp):
            key = f"{prefix}/{path.relative_to(out).as_posix()}"
            storage.put(key, path.read_bytes())
            artifacts[path.relative_to(out).as_posix()] = key
    manifest = {"prefix": prefix, "files": artifacts, "slide_count": len(draft.slide_content.slides), "media": media}
    storage.put_json(manifest_key, manifest)
    return manifest


def attach_notes(job_id, draft, rendered):
    from pptx import Presentation
    storage = Storage()
    raw = storage.get(rendered["files"]["slides.pptx"])
    prs = Presentation(io.BytesIO(raw))
    for slide, spec in zip(prs.slides, draft.slide_content.slides, strict=True):
        visual = rendered.get("media", {}).get(spec.id, {})
        slide.notes_slide.notes_text_frame.text = script_notes(spec) + "\n\n" + image_credit(visual) + "\n" + visual.get("source_url", "")
    buffer = io.BytesIO()
    prs.save(buffer)
    key = f"jobs/{job_id}/drafts/{digest([draft.model_dump(), rendered.get('media', {})])}/slides.pptx"
    if not storage.exists(key):
        storage.put(key, buffer.getvalue())
    return {**rendered, "files": {**rendered["files"], "slides.pptx": key}}


def export_docx(draft, source, dest, mode):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = settings().slide_font
    normal.font.size = Pt(11)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "MAIC · Giáo án và lời giảng" + (" · DEMO" if mode == "demo" else "")
    doc.add_heading(draft.lesson.title, 0)
    doc.add_paragraph("Đối tượng: " + draft.lesson.audience)
    doc.add_heading("Mục tiêu", 1)
    for objective in draft.lesson.objectives:
        doc.add_paragraph(objective.text, style="List Bullet")
    for title, items in [("Kiến thức cần có", draft.lesson.prerequisites), ("Chuẩn bị", draft.lesson.materials)]:
        doc.add_heading(title, 1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Tiến trình bài dạy", 1)
    for activity in draft.lesson.activities:
        doc.add_heading(f"{activity.title} · {activity.minutes} phút", 2)
        for label, value in [("Giáo viên", activity.teacher_action), ("Học viên", activity.learner_action), ("Đánh giá", activity.assessment)]:
            doc.add_paragraph(label + ": " + value)
        doc.add_paragraph("Nguồn: " + ", ".join(activity.evidence_ids))
    doc.add_heading("Bài tập sau buổi học", 1)
    doc.add_paragraph(draft.lesson.homework)
    doc.add_page_break()
    doc.add_heading("Script giảng theo slide", 1)
    for i, slide in enumerate(draft.slide_content.slides):
        doc.add_heading(f"Slide {i+1}. {slide.title}", 2)
        for para in script_notes(slide).split("\n\n"):
            doc.add_paragraph(para)
    doc.add_heading("Danh mục nguồn", 1)
    used = {e for slide in draft.slide_content.slides for e in slide.evidence_ids + slide.teaching_script.evidence_ids}
    used.update(e for item in draft.lesson.objectives + draft.lesson.activities for e in item.evidence_ids)
    used.update(e for slide in draft.slide_content.slides for q in slide.teaching_script.questions for e in q.evidence_ids)
    for unit in source.units:
        if unit.id in used:
            doc.add_paragraph(f"{unit.id} · Trang {unit.page}: {unit.content[:300]}")
    doc.save(dest)


def convert_docx_pdf(path: Path):
    if renderer_backend() == "microsoft-office":
        output = office_pdf(path, path.with_suffix(".pdf"))
        with fitz.open(output) as doc:
            if not len(doc):
                raise RuntimeError("Giáo án PDF rỗng")
        return output
    with tempfile.TemporaryDirectory(prefix="maic-lo-", dir=path.parent) as profile:
        env = os.environ.copy()
        env.pop("OPENAI_API_KEY", None)
        result = subprocess.run([soffice_path(), "-env:UserInstallation=" + Path(profile).as_uri(),
                                 "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)],
                                capture_output=True, timeout=180, check=False, env=env)
        output = path.with_suffix(".pdf")
        if result.returncode or not output.exists():
            raise RuntimeError("Không thể xuất giáo án PDF")
        with fitz.open(output) as doc:
            if not len(doc):
                raise RuntimeError("Giáo án PDF rỗng")
        return output


def export_speaker_handout(draft, rendered, dest):
    """Use flowing Word paragraphs so long notes are not clipped to a single notes page."""
    doc = Document()
    doc.styles["Normal"].font.name = settings().slide_font
    doc.styles["Normal"].font.size = Pt(11)
    storage = Storage()
    for i, slide in enumerate(draft.slide_content.slides):
        if i:
            doc.add_page_break()
        doc.add_heading(f"Slide {i+1}. {slide.title}", 1)
        key = rendered["files"][f"render/slide-{i+1}.png"]
        doc.add_picture(io.BytesIO(storage.get(key)), width=DocInches(6))
        visual = rendered.get("media", {}).get(slide.id, {})
        if visual.get("asset_key"):
            doc.add_paragraph(image_credit(visual) + "\n" + visual.get("source_url", ""))
        for paragraph in script_notes(slide).split("\n\n"):
            doc.add_paragraph(paragraph)
    doc.save(dest)


if __name__ == "__main__":
    spec = Path(sys.argv[1]).resolve()
    data = json.loads(spec.read_text(encoding="utf-8"))
    report = build_deck(Draft.model_validate(data["draft"]), SourceDocument.model_validate(data["source"]),
                        spec.parent, data["mode"], media=data.get("media", {}))
    (spec.parent / "geometry.json").write_text(json.dumps(report, ensure_ascii=False), encoding="utf-8")
